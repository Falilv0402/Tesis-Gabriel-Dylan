# -*- coding: utf-8 -*-
"""
Edita en sitio (preservando formato) el excel de Historias de Usuario:
1) Retira HU030 (descartada) y renumera HU031-036 -> HU030-035.
2) Ajusta la tabla EPICAS (columna D) para que siga apuntando a los IDs correctos.
3) Redefine la redacción de escenarios ambiguos para que coincidan con lo que
   el sistema realmente hace hoy (ver auditoria_datos.py para el respaldo).
Guarda como v1.3 en Downloads y también una copia en docs/excel del repo.
"""
import openpyxl

SRC = r"C:\Users\Usuario\Downloads\P20261012_Historias de Usuario y Criterios de Validacion v1.2.xlsx"
OUT_DOWNLOADS = r"C:\Users\Usuario\Downloads\P20261012_Historias de Usuario y Criterios de Validacion v1.3.xlsx"
OUT_REPO = r"C:\Users\Usuario\OneDrive\Escritorio\TesisDG-ML\PROYECTO-TESIS-DG\docs\excel\P20261012_Historias de Usuario y Criterios de Validacion v1.3.xlsx"

wb = openpyxl.load_workbook(SRC)
ws = wb["HU"]

# ── 1) Copiar el bloque HU031-036 (filas 68-79) dos filas hacia arriba (66-77),
#      pisando el bloque de HU030. Cada bloque tiene EXACTAMENTE 2 escenarios,
#      así que la forma de las celdas combinadas ya calza perfecto.
# Columnas B,C,D,E (hu_id/rol/funcionalidad/razon) están fusionadas de 2 en 2;
# solo la fila "ancla" (la primera de cada bloque) tiene un valor escribible.
anchor_rows = {68, 70, 72, 74, 76, 78}
for old_row in range(68, 80):
    new_row = old_row - 2
    for col in range(6, 11):  # F..J (no combinadas, varían por escenario)
        ws.cell(row=new_row, column=col).value = ws.cell(row=old_row, column=col).value
    if old_row in anchor_rows:
        for col in range(2, 6):  # B..E (ancla del bloque combinado)
            ws.cell(row=new_row, column=col).value = ws.cell(row=old_row, column=col).value

# Relabel de los IDs (ya se copió el texto viejo "HU031".."HU036"; lo corregimos)
for new_row, new_id in [(66, "HU030"), (68, "HU031"), (70, "HU032"),
                         (72, "HU033"), (74, "HU034"), (76, "HU035")]:
    ws.cell(row=new_row, column=2).value = new_id

# Limpiar las 2 filas finales, ya obsoletas (openpyxl NO reajusta bien los
# merges al usar delete_rows sobre celdas combinadas -- deja rangos fusionados
# huérfanos que corrompen las dimensiones del archivo al reabrirlo. Más seguro:
# desfusionar y vaciar en sitio, dejando 2 filas en blanco al final).
for rng in ("B78:B79", "C78:C79", "D78:D79", "E78:E79"):
    ws.unmerge_cells(rng)
for row in (78, 79):
    for col in range(2, 11):  # B..J (columna A es un espaciador decorativo fusionado A72:A79, no se toca)
        ws.cell(row=row, column=col).value = None

print("HU sheet: max_row ahora", ws.max_row)

# ── 2) EPICAS: achicar el bloque EP06 (5->4 HUs) y mover el bloque EP07 un lugar arriba
epicas_ws = wb["EPICAS"]
epicas_ws.unmerge_cells("B30:B34")
epicas_ws.unmerge_cells("C30:C34")
epicas_ws.unmerge_cells("B35:B38")
epicas_ws.unmerge_cells("C35:C38")

ep07_label = epicas_ws["B35"].value
ep07_obj = epicas_ws["C35"].value
epicas_ws["B35"].value = None
epicas_ws["C35"].value = None
epicas_ws["B34"].value = ep07_label
epicas_ws["C34"].value = ep07_obj

new_d_values = ["HU028", "HU029", "HU030", "HU031", "HU032", "HU033", "HU034", "HU035"]
for i, val in enumerate(new_d_values):
    epicas_ws.cell(row=30 + i, column=4).value = val
epicas_ws["D38"].value = None

epicas_ws.merge_cells("B30:B33")
epicas_ws.merge_cells("C30:C33")
epicas_ws.merge_cells("B34:B37")
epicas_ws.merge_cells("C34:C37")

# ── 3) Redefinir escenarios ambiguos (Criterio=G, Contexto=H, Evento=I, Resultado=J)
def set_escenario(row, criterio=None, contexto=None, evento=None, resultado=None):
    if criterio is not None:
        ws.cell(row=row, column=7).value = criterio
    if contexto is not None:
        ws.cell(row=row, column=8).value = contexto
    if evento is not None:
        ws.cell(row=row, column=9).value = evento
    if resultado is not None:
        ws.cell(row=row, column=10).value = resultado


# HU001 esc3 (CP003) — validación de a un mensaje a la vez, no simultánea
set_escenario(5,
    contexto="Campo requerido vacío, distrito no seleccionado, contraseña sin el formato exigido, o correo de un dominio no institucional (gmail, hotmail, etc.)",
    resultado="Sistema muestra el primer error de validación que encuentra (un mensaje a la vez, no todos los campos a la vez)")

# HU005 esc1 (CP011) — no hay flujo de invitación dedicado
set_escenario(13,
    resultado="Sistema registra al usuario con la contraseña que el administrador define en el propio formulario y lo asocia a su colegio/distrito; Supabase envía el correo de confirmación estándar al nuevo usuario")

# HU005 esc4 (CP014) — idem, para coordinador
set_escenario(16,
    resultado="Sistema crea la cuenta con la contraseña definida por el administrador, la asocia al colegio seleccionado, y Supabase envía el correo de confirmación estándar al coordinador")

# HU007 esc2 (CP018) — modo descriptivo real, no un mensaje literal "Datos insuficientes"
set_escenario(20,
    contexto="Colegio sin los 4 bimestres cargados todavía",
    evento="Accede al dashboard del colegio",
    resultado="Sistema muestra el aviso 'Este colegio todavía no tiene los 4 bimestres cargados — el modelo está operando con datos limitados' y opera en modo descriptivo (no predictivo) en vez de bloquear el análisis")

# HU007 esc3 (CP019) — error de conexión genérico, no un mensaje de modelo específico
set_escenario(21,
    contexto="Servidor de análisis (FastAPI) no responde o falla la conexión",
    resultado="Sistema muestra un aviso de error de conexión indicando que no se pudo procesar la solicitud")

# HU011 esc2 (CP027) — error técnico crudo, no traducido a lenguaje de negocio
set_escenario(29,
    resultado="Sistema muestra el mensaje de error técnico devuelto por el proceso de entrenamiento (sin traducir a lenguaje de negocio)")

# HU012 esc1 (CP028) — SHAP solo en EM2022; colegio propio usa desglose de notas como proxy
set_escenario(30,
    resultado="En el modelo nacional (EM2022) muestra los factores con mayor peso vía SHAP; en el modelo propio del colegio muestra el desglose de notas por área como proxy de los factores de riesgo")

# HU012 esc2 (CP029) — la predicción ya es automática, no hay paso manual previo
set_escenario(31,
    contexto="Análisis de factores del alumno aún calculándose",
    evento="Consulta los factores justo después de seleccionar al alumno",
    resultado="Sistema muestra 'Cargando análisis de factores...' mientras se calcula (la predicción ya se ejecuta automáticamente, no requiere un paso manual previo)")

# HU017 esc2 (CP039) — no existe búsqueda libre por ID que pueda fallar
set_escenario(41,
    criterio="Alumno con análisis aún cargando",
    contexto="Alumno recién seleccionado de la lista, análisis todavía calculándose",
    evento="Solicita el detalle de riesgo",
    resultado="Sistema muestra un estado de carga en vez de la explicación (los alumnos siempre se eligen de una lista ya cargada; no hay búsqueda libre por ID que pueda no existir)")

# HU020 esc2 (CP045) — las reglas están incorporadas en el sistema, siempre hay sugerencia
set_escenario(47,
    criterio="Recomendación siempre disponible",
    contexto="Estudiante identificado en riesgo",
    evento="Consulta al estudiante",
    resultado="Sistema siempre muestra una sugerencia (las reglas de recomendación están incorporadas en el sistema; no dependen de un catálogo externo que pueda estar vacío)")

# HU021 esc2 (CP047) — orden estable, sin segundo criterio todavía
set_escenario(49,
    criterio="Orden estable ante empate",
    contexto="Dos o más estudiantes con la misma probabilidad de riesgo",
    evento="Consulta la lista",
    resultado="Sistema mantiene el orden en que llegaron del análisis (aún no aplica un segundo criterio de desempate como promedio o asistencia)")

# HU022 esc1 (CP048) — segmentación por tipo solo en EM2022
set_escenario(50,
    resultado="En el modelo nacional (EM2022) agrupa por tipo específico (ej. 'Bajo en Lectura', 'Rendimiento múltiple'); en el modelo propio del colegio agrupa solo por nivel (ALTO/MEDIO/BAJO), sin tipo específico")

# HU022 esc2 (CP049) — reencuadre del límite real
set_escenario(51,
    contexto="Colegio con modelo propio (sin campo de tipo de riesgo disponible)",
    evento="Ejecuta el análisis",
    resultado="Sistema solo puede segmentar por nivel de riesgo, no por tipo específico")

# HU023 esc2 (CP051) — el botón queda deshabilitado, no hay mensaje de error posterior
set_escenario(53,
    criterio="Registro bloqueado sin alumno seleccionado",
    contexto="Ningún alumno seleccionado en el formulario",
    evento="Intenta guardar la intervención",
    resultado="Sistema mantiene el botón 'Registrar' deshabilitado hasta que se seleccione un alumno (el resto de campos ya tiene valores por defecto)")

# HU025 esc1 (CP054) — el histórico se exporta en CSV, no en PDF
set_escenario(56,
    resultado="Sistema exporta el historial de versiones del modelo en CSV (aparte, cada estudiante individual puede generar su propio reporte en PDF desde su vista de detalle)")

# HU026 esc2 (CP057) — no hay selector manual de periodos que pueda fallar
set_escenario(59,
    criterio="Comparación sobre todo el histórico disponible",
    contexto="Sin selector manual de periodos",
    evento="Accede al gráfico de histórico",
    resultado="Sistema muestra siempre la comparación de todo el histórico disponible (no hay un selector de periodo A-vs-B que pueda recibir una selección inválida)")

# HU027 esc2 (CP059) — exportación local, no depende de red, no hay modo de fallo real
set_escenario(61,
    criterio="Exportación con lista vacía",
    contexto="Ningún estudiante en la lista filtrada",
    evento="Solicita la exportación",
    resultado="Sistema genera igual el archivo, solo con encabezados (la exportación es una operación local que no depende de la red, por lo que no tiene un modo de fallo por conexión)")

# HU028 esc1 (CP060) — integración real es EM2022 + Excel del colegio, no APIs externas
set_escenario(62,
    contexto="Dataset nacional (EM2022) y notas internas del colegio (Excel) disponibles",
    evento="Selecciona el colegio en el dashboard",
    resultado="Sistema combina ambas fuentes en el mismo selector de colegios (no hay integración en vivo vía APIs externas — ambas se cargan como archivo)")

# HU028 esc2 (CP061) — error genérico de "backend desconectado", no identifica la fuente
set_escenario(63,
    contexto="Backend (FastAPI) desconectado",
    evento="Intenta cargar el dashboard",
    resultado="Sistema muestra un aviso genérico de 'backend desconectado' (no identifica por nombre cuál fuente específica falló)")

# ── Filas dentro del bloque que se recorrió (nuevas filas 69,70,71,73,74,77) ──

# HU031 nuevo, esc2 (CP067 nuevo, antes CP069) — ausencia de avisos = confirmación implícita
set_escenario(69,
    contexto="Datos limpios, sin advertencias durante la carga",
    resultado="Sistema no muestra ningún bloque de advertencias (la ausencia de avisos es la confirmación implícita de que todo se procesó sin errores)")

# HU032 nuevo, esc1 (CP068 nuevo, antes CP070) — se aplica de inmediato, no "al próximo entrenamiento"
set_escenario(70,
    resultado="Sistema aplica de inmediato los nuevos umbrales ALTO/MEDIO sobre las probabilidades ya calculadas (son umbrales de clasificación configurables, no hiperparámetros del algoritmo de entrenamiento)")

# HU032 nuevo, esc2 (CP069 nuevo, antes CP071) — prevención por rango fijo del control, no validación posterior
set_escenario(71,
    contexto="Control deslizante con rango mínimo/máximo fijo en la interfaz",
    resultado="Los controles deslizantes tienen un rango fijo; no es físicamente posible ingresar un valor fuera de rango")

# HU033 nuevo, esc2 (CP071 nuevo, antes CP073) — degrada a modo descriptivo, no bloquea
set_escenario(73,
    resultado="Sistema reentrena igual, pero cae a modo descriptivo (menos preciso) y lo advierte en el dashboard, en vez de bloquear el reentrenamiento")

# HU034 nuevo, esc1 (CP072 nuevo, antes CP074) — SHAP global solo en EM2022
set_escenario(74,
    resultado="En el modelo nacional (EM2022) muestra el ranking global de variables (SHAP); el modelo propio de cada colegio no lo expone, por el tamaño reducido de su muestra")

# HU035 nuevo, esc2 (CP075 nuevo, antes CP077) — solo expone fecha, no hay umbral de antigüedad
set_escenario(77,
    criterio="Fecha visible sin alerta de antigüedad",
    contexto="Modelo entrenado hace tiempo",
    resultado="Sistema siempre muestra la fecha exacta de la última actualización (aún no existe un umbral configurable de 'antigüedad máxima' que dispare una alerta automática)")

wb.save(OUT_DOWNLOADS)
wb.save(OUT_REPO)
print("guardado:", OUT_DOWNLOADS)
print("guardado:", OUT_REPO)
