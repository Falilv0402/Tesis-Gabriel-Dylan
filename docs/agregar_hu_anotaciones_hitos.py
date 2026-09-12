# -*- coding: utf-8 -*-
"""
agregar_hu_anotaciones_hitos.py -- Agrega las 2 HU base que faltaban desde
antes de esta sesión (pedido explícito de Mathias, 2026-09-12): "Anotaciones"
y "Plan de hitos" existen y funcionan en la app, pero nunca tuvieron su
propia HU/CP -- solo se documentó la aprobación (HU021 en v1.8) que se
construye ENCIMA de ellas.

Se insertan en EP04 (Priorización e Intervención), entre "Notificaciones de
equipo" y "Enviar plan a revisión" (orden narrativo: primero se anota/agrega
hitos, luego se envía a revisión). Todo lo demás se renumera automáticamente.

Parte de v1.8 (HU) / v1.7 (CP) y guarda como v1.9 / v1.8.
"""
import openpyxl

HU_SRC = r"C:\Users\Usuario\Downloads\P20261012_Historias de Usuario y Criterios de Validacion v1.8.xlsx"
HU_OUT_DL = r"C:\Users\Usuario\Downloads\P20261012_Historias de Usuario y Criterios de Validacion v1.9.xlsx"
HU_OUT_REPO = r"C:\Users\Usuario\OneDrive\Escritorio\TesisDG-ML\PROYECTO-TESIS-DG\docs\excel\P20261012_Historias de Usuario y Criterios de Validacion v1.9.xlsx"

CP_SRC = r"C:\Users\Usuario\Downloads\P20261012_Casos de Prueba v1.7.xlsx"
CP_OUT_DL = r"C:\Users\Usuario\Downloads\P20261012_Casos de Prueba v1.8.xlsx"
CP_OUT_REPO = r"C:\Users\Usuario\OneDrive\Escritorio\TesisDG-ML\PROYECTO-TESIS-DG\docs\excel\P20261012_Casos de Prueba v1.8.xlsx"

NEW_HUS = [
    {
        "key": "ANOTACIONES",
        "rol": "Director / Coordinador Académico",
        "quiere": "escribir anotaciones (notas de texto libre) sobre un alumno de mi colegio",
        "para": "dejar constancia de observaciones puntuales, visibles para todo el equipo",
        "esc": [
            {"num": 1, "nombre": "Anotación guardada",
             "dado": "Alumno seleccionado y texto no vacío",
             "cuando": "Escribe la anotación y presiona \"Guardar\"",
             "entonces": "Sistema la guarda, la muestra primero en la lista (más reciente arriba), notifica al equipo del colegio y la deja visible para cualquier usuario autenticado, no solo el autor"},
            {"num": 2, "nombre": "Sin anotaciones previas",
             "dado": "Alumno sin ninguna anotación registrada",
             "cuando": "Abre la pestaña \"Anotaciones\"",
             "entonces": "Sistema muestra \"Sin anotaciones para este alumno. Escribe la primera abajo.\""},
        ],
        "cps": [
            {"nombre": "Anotación guardada",
             "descripcion": "Validar que el sistema guarde la anotación y la muestre primero en la lista, visible para todo el equipo",
             "autor": "Director / Coordinador Académico",
             "precondiciones": "La aplicacion esta instalada y operativa. Alumno seleccionado",
             "pasos": [
                 ("Usuario abre la pestaña \"Anotaciones\" del alumno y escribe un texto", "Aplicación habilita el botón \"Guardar\""),
                 ("Usuario selecciona \"Guardar\"", "Aplicación guarda la anotación, la muestra primero en la lista y notifica al equipo del colegio"),
             ],
             "tipo": "Manual", "prioridad": "Media", "requerimientos": "Cuenta Director o Coordinador, alumno seleccionado",
             "postcondiciones": "Sistema deja la anotación visible para cualquier usuario autenticado, no solo el autor"},
            {"nombre": "Sin anotaciones previas",
             "descripcion": "Validar que el sistema muestre el mensaje de lista vacía cuando el alumno no tiene anotaciones",
             "autor": "Director / Coordinador Académico",
             "precondiciones": "La aplicacion esta instalada y operativa. Alumno sin ninguna anotación registrada",
             "pasos": [
                 ("Usuario abre la pestaña \"Anotaciones\" de un alumno sin anotaciones", "Aplicación muestra \"Sin anotaciones para este alumno. Escribe la primera abajo.\""),
             ],
             "tipo": "Automática", "prioridad": "Baja", "requerimientos": "Alumno sin anotaciones",
             "postcondiciones": "Sistema no bloquea la vista; invita a escribir la primera anotación"},
        ],
    },
    {
        "key": "PLAN_HITOS",
        "rol": "Director / Coordinador Académico",
        "quiere": "definir hitos concretos con fecha objetivo en el plan de seguimiento de un alumno",
        "para": "llevar un registro ordenado de las acciones planeadas para ese alumno",
        "esc": [
            {"num": 1, "nombre": "Hito agregado",
             "dado": "Un texto de hito y una fecha objetivo (con fecha por defecto si no se elige una)",
             "cuando": "Selecciona \"Agregar\"",
             "entonces": "Sistema lo guarda de inmediato y notifica al equipo del colegio; el hito queda visible para cualquier usuario autenticado, no solo el autor"},
            {"num": 2, "nombre": "Marcar como completado",
             "dado": "Un hito de cualquier miembro del equipo (el Director puede marcar cualquiera; el Coordinador solo los propios)",
             "cuando": "Selecciona la casilla del hito",
             "entonces": "Sistema lo marca como completado (tachado) si el usuario tiene permiso; si no, no guarda el cambio"},
        ],
        "cps": [
            {"nombre": "Hito agregado",
             "descripcion": "Validar que el sistema guarde el hito de inmediato y notifique al equipo del colegio",
             "autor": "Director / Coordinador Académico",
             "precondiciones": "La aplicacion esta instalada y operativa. Alumno seleccionado",
             "pasos": [
                 ("Usuario abre la pestaña \"Plan\" del alumno, escribe el texto del hito y elige (o deja) una fecha", "Aplicación habilita \"Agregar\""),
                 ("Usuario selecciona \"Agregar\"", "Aplicación guarda el hito de inmediato y notifica al equipo del colegio"),
             ],
             "tipo": "Manual", "prioridad": "Media", "requerimientos": "Cuenta Director o Coordinador, alumno seleccionado",
             "postcondiciones": "Sistema deja el hito visible para cualquier usuario autenticado, no solo el autor"},
            {"nombre": "Marcar como completado",
             "descripcion": "Validar que el Director pueda completar cualquier hito de su equipo, y el Coordinador solo los propios",
             "autor": "Director / Coordinador Académico",
             "precondiciones": "La aplicacion esta instalada y operativa. Hito sin completar de cualquier miembro del equipo",
             "pasos": [
                 ("Director selecciona la casilla de un hito de otro miembro del equipo", "Aplicación lo marca como completado (tachado)"),
                 ("Coordinador selecciona la casilla de un hito ajeno", "Aplicación no ofrece la casilla habilitada para ese caso (solo para los propios)"),
             ],
             "tipo": "Manual", "prioridad": "Media", "requerimientos": "Cuenta Director o Coordinador",
             "postcondiciones": "Sistema aplica el cambio de estado solo cuando el rol de quien lo hace tiene permiso"},
        ],
    },
]

# ═══════════════════════════════════════════════════════════════════════════
# 1) HU + EPICAS
# ═══════════════════════════════════════════════════════════════════════════
wb_hu = openpyxl.load_workbook(HU_SRC)
ws_hu = wb_hu["HU"]


def leer_hus(ws):
    records, row, cur = [], 3, None
    while row <= ws.max_row:
        id_val = ws.cell(row=row, column=2).value
        if id_val:
            cur = {"id": id_val, "rol": ws.cell(row=row, column=3).value,
                   "quiere": ws.cell(row=row, column=4).value,
                   "para": ws.cell(row=row, column=5).value, "esc": []}
            records.append(cur)
        esc_num = ws.cell(row=row, column=6).value
        if esc_num:
            cur["esc"].append({
                "num": esc_num, "nombre": ws.cell(row=row, column=7).value,
                "dado": ws.cell(row=row, column=8).value,
                "cuando": ws.cell(row=row, column=9).value,
                "entonces": ws.cell(row=row, column=10).value,
            })
        row += 1
    return records


hu_old = leer_hus(ws_hu)
assert len(hu_old) == 33, f"esperaba 33 HU en v1.8, hay {len(hu_old)}"
by_id = {r["id"]: r for r in hu_old}
HU_BY_KEY = {h["key"]: h for h in NEW_HUS}

epi = wb_hu["EPICAS"]
epicas_old, row = [], 3
while row <= epi.max_row:
    label = epi.cell(row=row, column=2).value
    if label:
        start = row
        obj = epi.cell(row=row, column=3).value
        hu_ids, r = [], row
        while r <= epi.max_row and (r == start or not epi.cell(row=r, column=2).value):
            hid = epi.cell(row=r, column=4).value
            if hid:
                hu_ids.append(hid)
            r += 1
        epicas_old.append({"label": label, "obj": obj, "hu_ids": hu_ids})
        row = r
    else:
        row += 1

assert sum(len(e["hu_ids"]) for e in epicas_old) == 33
ep04 = next(e for e in epicas_old if e["label"].startswith("EP04"))
# EP04 hoy: HU017,HU018,HU019,HU020,HU021(enviar plan a revisión) -- insertar
# ANOTACIONES y PLAN_HITOS antes de HU021 (justo después de notificaciones).
idx_hu021 = ep04["hu_ids"].index("HU021")
ep04["hu_ids"][idx_hu021:idx_hu021] = ["ANOTACIONES", "PLAN_HITOS"]

old_to_new_hu, final_hu_order, counter = {}, [], 1
for ep in epicas_old:
    for hid in ep["hu_ids"]:
        new_id = f"HU{counter:03d}"
        if hid in HU_BY_KEY:
            rec = HU_BY_KEY[hid]
        else:
            rec = by_id[hid]
            old_to_new_hu[hid] = new_id
        final_hu_order.append((new_id, rec))
        counter += 1

TOTAL_HU = counter - 1
assert TOTAL_HU == 35, f"esperaba 35 HU finales, salieron {TOTAL_HU}"
print(f"HU finales: {TOTAL_HU}")

for m in list(ws_hu.merged_cells.ranges):
    if m.min_row >= 3:
        ws_hu.unmerge_cells(str(m))
for row_ in ws_hu.iter_rows(min_row=3, max_row=ws_hu.max_row):
    for cell in row_:
        cell.value = None

row = 3
for new_id, rec in final_hu_order:
    n_esc = len(rec["esc"])
    start_row = row
    ws_hu.cell(row=start_row, column=2).value = new_id
    ws_hu.cell(row=start_row, column=3).value = rec["rol"]
    ws_hu.cell(row=start_row, column=4).value = rec["quiere"]
    ws_hu.cell(row=start_row, column=5).value = rec["para"]
    if n_esc > 1:
        end_row = start_row + n_esc - 1
        for col in (2, 3, 4, 5):
            ws_hu.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)
    for i, e in enumerate(rec["esc"]):
        r = start_row + i
        ws_hu.cell(row=r, column=6).value = e["num"]
        ws_hu.cell(row=r, column=7).value = e["nombre"]
        ws_hu.cell(row=r, column=8).value = e["dado"]
        ws_hu.cell(row=r, column=9).value = e["cuando"]
        ws_hu.cell(row=r, column=10).value = e["entonces"]
    row = start_row + n_esc
print(f"HU sheet: filas 3..{row - 1}")

for m in list(epi.merged_cells.ranges):
    epi.unmerge_cells(str(m))
for row_ in epi.iter_rows(min_row=3, max_row=epi.max_row):
    for cell in row_:
        cell.value = None

row = 3
key_to_new_id = {}
for new_id, rec in final_hu_order:
    for k, h in HU_BY_KEY.items():
        if h is rec:
            key_to_new_id[k] = new_id

for ep in epicas_old:
    n = len(ep["hu_ids"])
    start_row, end_row = row, row + n - 1
    epi.cell(row=start_row, column=2).value = ep["label"]
    epi.cell(row=start_row, column=3).value = ep["obj"]
    if n > 1:
        epi.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
        epi.merge_cells(start_row=start_row, start_column=3, end_row=end_row, end_column=3)
    for i, hid in enumerate(ep["hu_ids"]):
        epi.cell(row=start_row + i, column=4).value = old_to_new_hu.get(hid) or key_to_new_id.get(hid)
    row = end_row + 1
print(f"EPICAS: filas 3..{row - 1}")

wb_hu.save(HU_OUT_DL)
wb_hu.save(HU_OUT_REPO)
print("guardado:", HU_OUT_DL)

# ═══════════════════════════════════════════════════════════════════════════
# 2) LISTA CP + pestañas
# ═══════════════════════════════════════════════════════════════════════════
wb_cp = openpyxl.load_workbook(CP_SRC)
lst = wb_cp["LISTA CP"]

cps_by_old_hu, row = {}, 3
while row <= lst.max_row:
    cp_id = lst.cell(row=row, column=2).value
    if cp_id:
        hid = lst.cell(row=row, column=4).value
        cps_by_old_hu.setdefault(hid, []).append({
            "old_id": cp_id, "descripcion": lst.cell(row=row, column=3).value,
            "numero": lst.cell(row=row, column=5).value, "nombre": lst.cell(row=row, column=6).value,
        })
    row += 1

total_old_cp = sum(len(v) for v in cps_by_old_hu.values())
assert total_old_cp == 74, f"esperaba 74 CP en v1.7, hay {total_old_cp}"

actions, counter = [], 1
for hu_num in range(1, TOTAL_HU + 1):
    new_hu_id = f"HU{hu_num:03d}"
    old_hid = next((h for h, n in old_to_new_hu.items() if n == new_hu_id), None)
    if old_hid is None:
        key = next(k for k, nid in key_to_new_id.items() if nid == new_hu_id)
        for e in HU_BY_KEY[key]["cps"]:
            new_cp_id = f"CP{counter:03d}"
            actions.append({"new_cp_id": new_cp_id, "new_hu_id": new_hu_id, "kind": "NEW",
                             "numero": HU_BY_KEY[key]["cps"].index(e) + 1, "nombre": e["nombre"],
                             "descripcion": e["descripcion"], "content": e})
            counter += 1
    else:
        for cp in cps_by_old_hu[old_hid]:
            new_cp_id = f"CP{counter:03d}"
            actions.append({"new_cp_id": new_cp_id, "new_hu_id": new_hu_id, "kind": "COPY",
                             "old_id": cp["old_id"], "descripcion": cp["descripcion"],
                             "numero": cp["numero"], "nombre": cp["nombre"]})
            counter += 1

TOTAL_CP = counter - 1
print(f"CP finales: {TOTAL_CP}")
assert TOTAL_CP == 78, f"esperaba 78 CP finales, salieron {TOTAL_CP}"

for m in list(lst.merged_cells.ranges):
    if m.min_row >= 3:
        lst.unmerge_cells(str(m))
for row_ in lst.iter_rows(min_row=3, max_row=lst.max_row):
    for cell in row_:
        cell.value = None

for i, a in enumerate(actions):
    r = 3 + i
    lst.cell(row=r, column=2).value = a["new_cp_id"]
    lst.cell(row=r, column=3).value = a["descripcion"]
    lst.cell(row=r, column=4).value = a["new_hu_id"]
    lst.cell(row=r, column=5).value = a["numero"]
    lst.cell(row=r, column=6).value = a["nombre"]

import re
copy_actions = [a for a in actions if a["kind"] == "COPY"]
new_actions = [a for a in actions if a["kind"] == "NEW"]

for a in copy_actions:
    tmp_ws = wb_cp.copy_worksheet(wb_cp[a["old_id"]])
    tmp_ws.title = f"TMP_{a['old_id']}"
for name in list(wb_cp.sheetnames):
    if re.fullmatch(r"CP\d{3}", name):
        del wb_cp[name]
for a in copy_actions:
    ws = wb_cp[f"TMP_{a['old_id']}"]
    ws.title = a["new_cp_id"]
    old_prefix = f"Caso de Prueba: {a['old_id']}:"
    new_prefix = f"Caso de Prueba: {a['new_cp_id']}:"
    a1 = ws["A1"].value or ""
    assert a1.startswith(old_prefix), f"{a['old_id']}: encabezado inesperado: {a1!r}"
    ws["A1"].value = new_prefix + a1[len(old_prefix):]


def crear_pestana_cp(wb, nombre, contenido):
    ws_cp = wb.create_sheet(nombre)
    ws_cp["A1"] = f"Caso de Prueba: {nombre}: {contenido['nombre']}"
    ws_cp["A2"] = "Autor:"; ws_cp["B2"] = contenido["autor"]
    ws_cp["A3"] = f"Precondiciones: {contenido['precondiciones']}"
    ws_cp["A6"] = "#:"; ws_cp["B6"] = "Pasos:"; ws_cp["C6"] = "Resultados Esperados:"
    r = 7
    for i, (paso, resultado) in enumerate(contenido["pasos"], 1):
        ws_cp.cell(row=r, column=1).value = i
        ws_cp.cell(row=r, column=2).value = paso
        ws_cp.cell(row=r, column=3).value = resultado
        r += 1
    ws_cp.cell(row=r, column=1).value = "Tipo de ejecución:"; ws_cp.cell(row=r, column=2).value = contenido["tipo"]; r += 1
    ws_cp.cell(row=r, column=1).value = "Prioridad:"; ws_cp.cell(row=r, column=2).value = contenido["prioridad"]; r += 1
    ws_cp.cell(row=r, column=1).value = "Requerimientos"; ws_cp.cell(row=r, column=2).value = contenido["requerimientos"]; r += 1
    ws_cp.cell(row=r, column=1).value = f"Postcondiciones:\n{contenido['postcondiciones']}"
    return ws_cp


for a in new_actions:
    crear_pestana_cp(wb_cp, a["new_cp_id"], a["content"])

cp_sheet_names = sorted(n for n in wb_cp.sheetnames if re.fullmatch(r"CP\d{3}", n))
expected_names = sorted(a["new_cp_id"] for a in actions)
assert cp_sheet_names == expected_names
assert not any(n.startswith("TMP_") for n in wb_cp.sheetnames)

hu_esc_count = {new_id: len(rec["esc"]) for new_id, rec in final_hu_order}
cp_count_by_hu = {}
for a in actions:
    cp_count_by_hu[a["new_hu_id"]] = cp_count_by_hu.get(a["new_hu_id"], 0) + 1
mismatches = [(hu, hu_esc_count[hu], cp_count_by_hu.get(hu, 0)) for hu in hu_esc_count if hu_esc_count[hu] != cp_count_by_hu.get(hu, 0)]
assert not mismatches, f"desbalance HU/CP: {mismatches}"

wb_cp.save(CP_OUT_DL)
wb_cp.save(CP_OUT_REPO)
print("guardado:", CP_OUT_DL)
print(f"Pestañas CP finales: {len(cp_sheet_names)}")
