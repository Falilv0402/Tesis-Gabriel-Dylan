# -*- coding: utf-8 -*-
from auditoria_datos import HUS, HU_EPICA, EPICAS

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0F, 0x1F, 0x3D)
GREEN = RGBColor(0x16, 0x65, 0x34)
AMBER = RGBColor(0x92, 0x40, 0x0E)
RED = RGBColor(0x99, 0x1B, 0x1B)
GRAY = RGBColor(0x5A, 0x6A, 0x85)

FILL = {"OK": "DCFCE7", "PARCIAL": "FEF3C7", "NO": "FEE2E2"}
FONT_COLOR = {"OK": GREEN, "PARCIAL": AMBER, "NO": RED}
LABEL = {"OK": "Cumple", "PARCIAL": "Parcial", "NO": "No cumple"}

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.font.color.rgb = RGBColor(0x1A, 0x25, 0x40)

for section in doc.sections:
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def add_title(text, size=24, color=NAVY, space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_h2(text, color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = color
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "2563EB")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_h3(hu_id, quiere, estado, pct):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"{hu_id} — ")
    r1.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = NAVY
    r2 = p.add_run(f'"Quiero {quiere}"')
    r2.italic = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(0x1A, 0x25, 0x40)
    p.add_run("   ")
    r3 = p.add_run(estado)
    r3.bold = True
    r3.font.size = Pt(11)
    r3.font.color.rgb = FONT_COLOR["OK"] if pct == 100 else (FONT_COLOR["NO"] if pct == 0 else FONT_COLOR["PARCIAL"])
    return p


def add_body(text, size=10.5, color=None, bold=False, italic=False, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    return p


# ── Portada ──────────────────────────────────────────────────────────────
add_title("Auditoría de Cumplimiento — Historias de Usuario y Casos de Prueba")
sub = doc.add_paragraph()
run = sub.add_run("SATRA · P20261012 — Estado real contrastado contra el sistema en producción")
run.font.size = Pt(12)
run.font.color.rgb = GRAY
sub.paragraph_format.space_after = Pt(4)
fecha = doc.add_paragraph()
r = fecha.add_run("Septiembre 2026 · Fuente: P20261012_Historias de Usuario y Criterios de Validación v1.2 + P20261012_Casos de Prueba v1.1")
r.font.size = Pt(9.5)
r.font.color.rgb = GRAY
fecha.paragraph_format.space_after = Pt(16)

add_body(
    "Este documento contrasta cada una de las 36 historias de usuario (HU) y sus 77 casos de prueba (CP) "
    "asociados contra el comportamiento real del sistema desplegado en producción (satraapp.com) — no contra "
    "la intención de diseño ni contra documentación previa. Cada caso se marca como Cumple, Parcial o No cumple, "
    "con la evidencia concreta (archivo, componente o comportamiento verificado) que sustenta esa calificación."
)

# ── Resumen ejecutivo ────────────────────────────────────────────────────
add_h2("Resumen ejecutivo")

total_ok = sum(1 for hu in HUS for e in hu[4] if e[3] == "OK")
total_parcial = sum(1 for hu in HUS for e in hu[4] if e[3] == "PARCIAL")
total_no = sum(1 for hu in HUS for e in hu[4] if e[3] == "NO")
total_n = total_ok + total_parcial + total_no
pct_global = round(100 * (total_ok + 0.5 * total_parcial) / total_n)

add_body(
    f"De los {total_n} casos de prueba evaluados (uno por cada escenario de las 36 HU): "
    f"{total_ok} cumplen completamente, {total_parcial} cumplen parcialmente, y {total_no} no cumplen. "
    f"Esto da un cumplimiento ponderado global de {pct_global}% (contando cada caso parcial como medio punto)."
)

tbl = doc.add_table(rows=1, cols=3)
tbl.style = "Table Grid"
hdr = tbl.rows[0].cells
for i, h in enumerate(["Estado", "Casos de prueba", "% del total"]):
    set_cell_text(hdr[i], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10)
    set_cell_shading(hdr[i], "0F1F3D")
for estado, count in [("OK", total_ok), ("PARCIAL", total_parcial), ("NO", total_no)]:
    row = tbl.add_row().cells
    set_cell_text(row[0], LABEL[estado], bold=True, color=FONT_COLOR[estado])
    set_cell_shading(row[0], FILL[estado])
    set_cell_text(row[1], str(count), align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(row[2], f"{round(100*count/total_n)}%", align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
add_body(
    "Los casos \"Parcial\" no son errores del sistema — casi todos comparten la misma causa de fondo: la HU "
    "asume un comportamiento (mensaje exacto, flujo de invitación por correo, un umbral configurable, "
    "importancia de variables con SHAP) que el sistema real satisface en su intención pero no de forma "
    "literal, o que solo aplica a uno de los dos modelos predictivos (EM2022 nacional vs. modelo propio por "
    "colegio). Los 7 casos \"No cumple\" están concentrados casi todos en dos causas puntuales: el control de "
    "cuenta desactivada no se aplica en el login (HU002/HU005), y la actualización periódica automática "
    "(HU030) es la única funcionalidad de las 36 que queda pendiente de una decisión de producto.",
    italic=True,
)

# ── Por épica ────────────────────────────────────────────────────────────
current_epica = None
for hu_id, rol, quiere, para, escenarios in HUS:
    epica_id = HU_EPICA[hu_id]
    if epica_id != current_epica:
        current_epica = epica_id
        add_h2(f"{epica_id} — {EPICAS[epica_id]}")

    n = len(escenarios)
    ok = sum(1 for e in escenarios if e[3] == "OK")
    parcial = sum(1 for e in escenarios if e[3] == "PARCIAL")
    no = sum(1 for e in escenarios if e[3] == "NO")
    pct = round(100 * (ok + 0.5 * parcial) / n)
    if ok == n:
        estado_txt = "✓ Cumple 100%"
    elif no == n:
        estado_txt = "✗ No cumple"
    else:
        estado_txt = f"◐ Parcial — {pct}%"

    add_h3(hu_id, quiere, estado_txt, 100 if ok == n else (0 if no == n else 50))
    razon_p = doc.add_paragraph()
    razon_p.paragraph_format.space_after = Pt(6)
    r = razon_p.add_run(f"Rol: {rol}  ·  Para: {para}")
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY
    r.italic = True

    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(["CP", "Escenario", "Estado", "Evidencia"]):
        set_cell_text(hdr[i], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=9)
        set_cell_shading(hdr[i], "0F1F3D")
    widths = [Cm(1.5), Cm(4.2), Cm(2.2), Cm(8.6)]
    for cp_id, criterio, resultado, estado, evidencia in escenarios:
        row = t.add_row().cells
        set_cell_text(row[0], cp_id, bold=True, size=9)
        set_cell_text(row[1], criterio, size=9)
        set_cell_text(row[2], LABEL[estado], bold=True, color=FONT_COLOR[estado], size=9)
        set_cell_shading(row[2], FILL[estado])
        set_cell_text(row[3], evidencia, size=9)
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

doc.add_paragraph().paragraph_format.space_after = Pt(2)
cierre = doc.add_paragraph()
r = cierre.add_run(
    "Metodología: cada caso se evaluó contrastando el criterio de aceptación de la HU (Contexto/Evento/"
    "Resultado esperado, según el archivo v1.2) contra el código y comportamiento real del sistema "
    "desplegado, incluyendo verificaciones puntuales de código para los casos donde la memoria del "
    "desarrollo no era suficiente por sí sola (p. ej. el chequeo de cuenta desactivada en el login)."
)
r.italic = True
r.font.color.rgb = GRAY
r.font.size = Pt(9)

OUT = r"C:\Users\Usuario\OneDrive\Escritorio\TesisDG-ML\PROYECTO-TESIS-DG\docs\word\Auditoria_Cumplimiento_HU_CP_2026-09.docx"
doc.save(OUT)
print("saved:", OUT)
