# -*- coding: utf-8 -*-
from auditoria_datos import HUS, HU_EPICA, EPICAS

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "0F1F3D"
WHITE = "FFFFFF"
LIGHTGREEN = "DCFCE7"
LIGHTAMBER = "FEF3C7"
LIGHTRED = "FEE2E2"

wb = Workbook()

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=WHITE, size=10.5)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="9CA3AF"))

def estado_fill(estado):
    return {"OK": LIGHTGREEN, "PARCIAL": LIGHTAMBER, "NO": LIGHTRED}.get(estado, "FFFFFF")

def estado_label(estado):
    return {"OK": "Cumple", "PARCIAL": "Parcial", "NO": "No cumple"}.get(estado, estado)

# ══════════════════════════════════════════════════════════════════════════
# Hoja 1: Resumen por HU
# ══════════════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Resumen por HU"
headers1 = ["HU", "Épica", "Rol", "Quiere...", "Total CP", "Cumple", "Parcial", "No cumple", "% Cumplimiento", "Estado"]
for i, h in enumerate(headers1, 1):
    ws1.cell(row=1, column=i, value=h)
style_header(ws1, 1, len(headers1))
ws1.row_dimensions[1].height = 30
ws1.freeze_panes = "A2"

total_ok = total_parcial = total_no = 0
r = 2
for hu_id, rol, quiere, para, escenarios in HUS:
    n = len(escenarios)
    ok = sum(1 for e in escenarios if e[3] == "OK")
    parcial = sum(1 for e in escenarios if e[3] == "PARCIAL")
    no = sum(1 for e in escenarios if e[3] == "NO")
    total_ok += ok; total_parcial += parcial; total_no += no
    pct = round(100 * (ok + 0.5 * parcial) / n)
    if ok == n:
        estado = "Cumple 100%"
    elif no == n:
        estado = "No cumple"
    else:
        estado = f"Parcial ({pct}%)"

    epica_id = HU_EPICA[hu_id]
    ws1.cell(row=r, column=1, value=hu_id).font = Font(bold=True)
    ws1.cell(row=r, column=2, value=f"{epica_id} — {EPICAS[epica_id]}")
    ws1.cell(row=r, column=3, value=rol)
    ws1.cell(row=r, column=4, value=quiere)
    ws1.cell(row=r, column=5, value=n)
    ws1.cell(row=r, column=6, value=ok)
    ws1.cell(row=r, column=7, value=parcial)
    ws1.cell(row=r, column=8, value=no)
    ws1.cell(row=r, column=9, value=pct / 100).number_format = "0%"
    estado_cell = ws1.cell(row=r, column=10, value=estado)
    fill_estado = "OK" if ok == n else ("NO" if no == n else "PARCIAL")
    estado_cell.fill = PatternFill("solid", fgColor=estado_fill(fill_estado))
    estado_cell.font = Font(bold=True)
    for c in range(1, len(headers1) + 1):
        ws1.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=(c in (2, 3, 4)))
    r += 1

# Fila de totales generales
r += 1
ws1.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws1.cell(row=r, column=5, value=total_ok + total_parcial + total_no).font = Font(bold=True)
ws1.cell(row=r, column=6, value=total_ok).font = Font(bold=True)
ws1.cell(row=r, column=7, value=total_parcial).font = Font(bold=True)
ws1.cell(row=r, column=8, value=total_no).font = Font(bold=True)
total_n = total_ok + total_parcial + total_no
ws1.cell(row=r, column=9, value=round(100 * (total_ok + 0.5 * total_parcial) / total_n) / 100).number_format = "0%"
ws1.cell(row=r, column=9).font = Font(bold=True)

widths1 = [8, 34, 22, 42, 9, 8, 8, 10, 13, 16]
for i, w in enumerate(widths1, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

# ══════════════════════════════════════════════════════════════════════════
# Hoja 2: Detalle por Caso de Prueba (77 CP)
# ══════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Detalle por CP")
headers2 = ["CP", "HU", "Épica", "Escenario / Criterio", "Resultado esperado (según HU)", "Estado", "Evidencia / Nota"]
for i, h in enumerate(headers2, 1):
    ws2.cell(row=1, column=i, value=h)
style_header(ws2, 1, len(headers2))
ws2.row_dimensions[1].height = 30
ws2.freeze_panes = "A2"

r = 2
for hu_id, rol, quiere, para, escenarios in HUS:
    epica_id = HU_EPICA[hu_id]
    for cp_id, criterio, resultado, estado, evidencia in escenarios:
        ws2.cell(row=r, column=1, value=cp_id).font = Font(bold=True)
        ws2.cell(row=r, column=2, value=hu_id)
        ws2.cell(row=r, column=3, value=epica_id)
        ws2.cell(row=r, column=4, value=criterio)
        ws2.cell(row=r, column=5, value=resultado)
        estado_cell = ws2.cell(row=r, column=6, value=estado_label(estado))
        estado_cell.fill = PatternFill("solid", fgColor=estado_fill(estado))
        estado_cell.font = Font(bold=True)
        ws2.cell(row=r, column=7, value=evidencia)
        for c in range(1, len(headers2) + 1):
            ws2.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=(c in (4, 5, 7)))
        r += 1

widths2 = [8, 8, 8, 26, 34, 12, 55]
for i, w in enumerate(widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

OUT = r"C:\Users\Usuario\OneDrive\Escritorio\TesisDG-ML\PROYECTO-TESIS-DG\docs\excel\Auditoria_Cumplimiento_HU_CP_2026-09.xlsx"
wb.save(OUT)
print("saved:", OUT)
