# -*- coding: utf-8 -*-
"""
corregir_incidencias_hu_cp.py
Corrige las 10 incidencias reportadas (de 18 totales, 1-10/18) sobre el
excel de Historias de Usuario y el de Casos de Prueba:

HU:
  - HU018: el criterio de aceptación no definía el texto de regla fija
    según nivel/tipo de riesgo (solo decía "un texto variable"). Se enumera
    la regla completa, igual a recommendation() en frontend/src/lib/format.ts.
  - HU019: el escenario "Registro exitoso" mezclaba dos comportamientos
    distintos (descripción con texto / descripción vacía) en un solo
    resultado, y referenciaba mal la regla fija (decía HU020, es HU018).
  - HU021 y HU022: el resultado documentaba la fuga de RLS que existía
    antes de la migración 0019 ("visible para cualquier usuario
    autenticado") en vez del alcance real y ya corregido (visible solo
    para el equipo del mismo colegio). HU022 no estaba en las 10
    incidencias reportadas pero tiene exactamente el mismo defecto que
    HU021 -- se corrige de paso.

CP (ambos defectos de fondo, presentes en todo el documento, no solo en
los 10 casos reportados -- se corrige donde corresponde):
  - "Autor" tenía el rol bajo prueba (ej. "Coordinador Académico") en
    vez de quién redactó el caso de prueba. Se reemplaza por "Gabriel
    Torres" en las 78 pestañas de CP.
  - CP024/028/029/030: precondiciones sin datos concretos/reproducibles
    (alumno o colegio genérico en vez de un ejemplo específico).
  - CP025: no se definía cómo reproducir el error de conexión.
  - CP026: los pasos describían un flujo de "Predecir alumnos en riesgo"
    con formulario de filtros y "obtener respuesta" que no corresponde
    al flujo real de HU010-1 (el dashboard muestra los niveles de riesgo
    directamente, sin ese paso manual).
  - CP042/CP044 (LISTA CP) y CP048/CP050 (postcondiciones): se actualizan
    para reflejar los mismos fixes de HU018/HU019/HU021/HU022 de arriba.
"""
import shutil
from pathlib import Path

import openpyxl

DOCS = Path(__file__).resolve().parent
EXCEL_DIR = DOCS / "excel"

HU_SRC = EXCEL_DIR / "P20261012_Historias de Usuario y Criterios de Validacion v1.9.xlsx"
HU_DST = EXCEL_DIR / "P20261012_Historias de Usuario y Criterios de Validacion v2.0.xlsx"
CP_SRC = EXCEL_DIR / "P20261012_Casos de Prueba v1.8.xlsx"
CP_DST = EXCEL_DIR / "P20261012_Casos de Prueba v1.9.xlsx"

AUTOR_QA = "Gabriel Torres"

RULE_TEXT_HU018 = (
    'Sistema guarda como descripción un texto de regla fija según el riesgo: '
    'BAJO → "Mantener seguimiento regular y revisar evolución en el siguiente periodo"; '
    'MEDIO → "Programar monitoreo académico y refuerzo preventivo"; '
    'ALTO con un área de bajo rendimiento identificada (Lectura, Ciencias, '
    'Matemática, Comunicación, varias áreas, rendimiento múltiple o contexto '
    'socioeconómico) → un texto de refuerzo propio de esa área; '
    'ALTO sin área específica identificada → "Monitoreo académico continuo y '
    'revisión periódica"'
)

RESULT_TEXT_HU019 = (
    'Sistema guarda la intervención con esa descripción (si se dejó vacía, ya '
    'llega resuelta con el texto de regla fija de HU018 -- no es un '
    'comportamiento propio de este escenario), estado inicial "pendiente" '
    'y fecha actual'
)


def visibilidad_equipo(sujeto: str) -> str:
    return (
        f'la deja visible para todo el equipo de ese colegio (cualquier '
        f'Director o Coordinador de esa IE, o superadmin), no para usuarios '
        f'de otros colegios'
    )


def find_row(ws, col, predicate, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, col).value
        if v is not None and predicate(str(v)):
            return r
    raise ValueError(f"No se encontró fila en {ws.title} con predicado dado")


# ============================================================
# 1) HU: corregir HU018/HU019/HU021/HU022
# ============================================================
shutil.copy(HU_SRC, HU_DST)
wb_hu = openpyxl.load_workbook(HU_DST)
ws_hu = wb_hu["HU"]

COL_ID, COL_ESC, COL_RES = 2, 6, 10  # B, F, J (1-indexed)

last_id = None
fixed_hu = {}
for r in range(3, ws_hu.max_row + 1):
    hid = ws_hu.cell(r, COL_ID).value
    if hid:
        last_id = hid
    esc = ws_hu.cell(r, COL_ESC).value

    if last_id == "HU018" and esc == 1:
        ws_hu.cell(r, COL_RES).value = RULE_TEXT_HU018
        fixed_hu["HU018-1"] = r
    elif last_id == "HU019" and esc == 1:
        ws_hu.cell(r, COL_RES).value = RESULT_TEXT_HU019
        fixed_hu["HU019-1"] = r
    elif last_id == "HU021" and esc == 1:
        old = ws_hu.cell(r, COL_RES).value
        new = old.split("y la deja visible")[0] + "y " + visibilidad_equipo("anotación")
        ws_hu.cell(r, COL_RES).value = new
        fixed_hu["HU021-1"] = r
    elif last_id == "HU022" and esc == 1:
        old = ws_hu.cell(r, COL_RES).value
        new = old.split("; el hito")[0] + "; el hito queda " + visibilidad_equipo("hito")
        ws_hu.cell(r, COL_RES).value = new
        fixed_hu["HU022-1"] = r

assert len(fixed_hu) == 4, f"Se esperaban 4 filas corregidas en HU, se corrigieron {len(fixed_hu)}: {fixed_hu}"
wb_hu.save(HU_DST)
print("HU corregido:", fixed_hu)


# ============================================================
# 2) CP: LISTA CP -- actualizar descripciones de CP042 y CP044
# ============================================================
shutil.copy(CP_SRC, CP_DST)
wb_cp = openpyxl.load_workbook(CP_DST)
ws_lista = wb_cp["LISTA CP"]

DESC_CP042 = (
    "Validar que el sistema guarde como descripción el texto de regla fija "
    "correspondiente al riesgo del alumno (BAJO, MEDIO, o ALTO con o sin "
    "área específica identificada — ver HU018)"
)
DESC_CP044 = (
    'Validar que el sistema guarde la intervención con la descripción '
    'escrita, estado inicial "pendiente" y fecha actual'
)

lista_fixed = []
for r in range(3, ws_lista.max_row + 1):
    cp_id = ws_lista.cell(r, 2).value
    if cp_id == "CP042":
        ws_lista.cell(r, 3).value = DESC_CP042
        lista_fixed.append("CP042")
    elif cp_id == "CP044":
        ws_lista.cell(r, 3).value = DESC_CP044
        lista_fixed.append("CP044")
assert set(lista_fixed) == {"CP042", "CP044"}, lista_fixed
print("LISTA CP corregida:", lista_fixed)


# ============================================================
# 3) CP: Autor -> "Gabriel Torres" en TODAS las pestañas CP0xx
# ============================================================
cp_tabs = [n for n in wb_cp.sheetnames if n.startswith("CP") and n != "LISTA CP"]
autor_fixed = []
for name in cp_tabs:
    ws = wb_cp[name]
    r = find_row(ws, 1, lambda v: v.strip() == "Autor:")
    ws.cell(r, 2).value = AUTOR_QA
    autor_fixed.append(name)
assert len(autor_fixed) == 78, f"Se esperaban 78 pestañas CP, se encontraron {len(autor_fixed)}"
print(f"Autor corregido en {len(autor_fixed)} pestañas CP")


# ============================================================
# 4) CP: precondiciones concretas (CP024, CP025, CP028, CP029, CP030)
# ============================================================
def set_precondicion(ws, nueva_precondicion_completa):
    r = find_row(ws, 1, lambda v: v.startswith("Precondiciones:"))
    ws.cell(r, 1).value = nueva_precondicion_completa


set_precondicion(
    wb_cp["CP024"],
    "Precondiciones: La aplicacion esta instalada y operativa. Se esta "
    "loggeado como Coordinador. Colegio de prueba con el Excel de notas "
    "cargado solo hasta Bimestre 3 (sin datos reales de Bimestre 4)",
)
set_precondicion(
    wb_cp["CP025"],
    "Precondiciones: La aplicacion esta instalada y operativa. Se esta "
    'loggeado como Coordinador. Backend FastAPI detenido (ej. "docker '
    'compose stop backend" en el servidor, o el proceso local '
    "interrumpido) mientras el frontend intenta cargar el dashboard",
)
set_precondicion(
    wb_cp["CP028"],
    "Precondiciones: La aplicacion esta instalada y operativa. Se esta "
    "loggeado como Coordinador. Alumno con nota registrada en al menos 1 "
    "de las 7 areas academicas + Conducta (ej. cualquier alumno de un "
    "colegio con modelo propio entrenado, como IE 831305)",
)
set_precondicion(
    wb_cp["CP029"],
    "Precondiciones: La aplicacion esta instalada y operativa. Se esta "
    "loggeado como Coordinador. Alumno del padron del colegio que no "
    "aparece en el Excel de notas subido (0 registros en las 7 areas + "
    "Conducta)",
)
set_precondicion(
    wb_cp["CP030"],
    "Precondiciones: La aplicacion esta instalada y operativa. Se esta "
    "loggeado como Coordinador. Colegio con modelo propio entrenado (ej. "
    "IE 831305, 349 alumnos) con un filtro de nivel, grado o seccion "
    "aplicado",
)
print("Precondiciones concretas: CP024, CP025, CP028, CP029, CP030")


# ============================================================
# 5) CP026: reescribir los Pasos para que correspondan a HU010-1
# ============================================================
ws026 = wb_cp["CP026"]
header_row = find_row(ws026, 1, lambda v: v.strip() == "#:")
nuevos_pasos = [
    (1, "Usuario esta loggeado como Coordinador", "Aplicacion muestra el dashboard"),
    (
        2,
        'Usuario accede al panel de niveles de riesgo (pestana "Estudiante")',
        "Aplicacion muestra la lista de alumnos con su nivel de riesgo "
        "(ALTO/MEDIO/BAJO) ya calculado, sin un paso manual de "
        '"predecir"',
    ),
    (
        3,
        "Usuario aplica un filtro de nivel, grado o seccion (opcional)",
        "Aplicacion refina la lista mostrada al instante, sin recargar",
    ),
]
for i, (num, paso, resultado) in enumerate(nuevos_pasos):
    r = header_row + 1 + i
    ws026.cell(r, 1).value = num
    ws026.cell(r, 2).value = paso
    ws026.cell(r, 3).value = resultado
# el paso 4 viejo ya no aplica -- limpiar esa fila
r4 = header_row + 1 + len(nuevos_pasos)
if ws026.cell(r4, 1).value == 4:
    ws026.cell(r4, 1).value = None
    ws026.cell(r4, 2).value = None
    ws026.cell(r4, 3).value = None
print("CP026: pasos reescritos para coincidir con HU010-1")


# ============================================================
# 6) CP042 / CP044: precondicion + resultado concretos
# ============================================================
ws042 = wb_cp["CP042"]
set_precondicion(
    ws042,
    "Precondiciones: La aplicacion esta instalada y operativa. Se esta "
    'loggeado como Coordinador. Alumno con nivel de riesgo ALTO y tipo '
    'de riesgo "Bajo en Matematica" (colegio propio) seleccionado, '
    "descripcion vacia",
)
header042 = find_row(ws042, 1, lambda v: v.strip() == "#:")
ws042.cell(header042 + 2, 3).value = (
    'Aplicacion guarda como descripcion "Refuerzo en Matematica: practica '
    'guiada y seguimiento semanal del progreso"'
)

ws044 = wb_cp["CP044"]
set_precondicion(
    ws044,
    "Precondiciones: La aplicacion esta instalada y operativa. Se esta "
    "loggeado como Coordinador. Alumno seleccionado, tipo de "
    "intervencion elegido y descripcion escrita",
)
header044 = find_row(ws044, 1, lambda v: v.strip() == "#:")
ws044.cell(header044 + 2, 3).value = (
    'Aplicacion guarda la intervencion con esa descripcion, estado '
    '"pendiente" y fecha actual'
)
print("CP042 y CP044: precondicion + resultado concretos")


# ============================================================
# 7) CP048 / CP050: postcondiciones -- mismo fix que HU021/HU022
# ============================================================
def set_postcondicion(ws, nueva):
    r = find_row(ws, 1, lambda v: v.startswith("Postcondiciones:"))
    ws.cell(r, 1).value = nueva


set_postcondicion(
    wb_cp["CP048"],
    "Postcondiciones:\nSistema deja la anotacion visible para todo el "
    "equipo de ese colegio, no para usuarios de otros colegios",
)
set_postcondicion(
    wb_cp["CP050"],
    "Postcondiciones:\nSistema deja el hito visible para todo el equipo "
    "de ese colegio, no para usuarios de otros colegios",
)
print("CP048 y CP050: postcondiciones corregidas")

wb_cp.save(CP_DST)
print("\nGuardado:", HU_DST.name, "y", CP_DST.name)
