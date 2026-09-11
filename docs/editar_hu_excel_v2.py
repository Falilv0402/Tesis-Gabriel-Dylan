# -*- coding: utf-8 -*-
"""
Segunda ronda de correcciones sobre el excel de Historias de Usuario, ahora
que EM2022 está apagado (EM2022_HABILITADO=false en el frontend):

1) Los 4 escenarios que todavía describían el modelo nacional (EM2022) se
   redefinen para describir SOLO lo que existe hoy en el modelo propio de
   cada colegio (que es lo único que un usuario real puede ver en la app).
2) HU022-esc2 se rediseña: con esc1 ya reducido a "solo por nivel", el
   escenario negativo más honesto y real no es "datos incompletos" (ya
   cubierto por esc1) sino "el colegio todavía no tiene modelo propio" --
   el mismo estado SinModeloPropio que ya existe de verdad en la app.

Parte de v1.3 (no del original del cliente) y guarda como v1.4.
"""
import openpyxl

SRC = r"C:\Users\Usuario\Downloads\P20261012_Historias de Usuario y Criterios de Validacion v1.3.xlsx"
OUT_DOWNLOADS = r"C:\Users\Usuario\Downloads\P20261012_Historias de Usuario y Criterios de Validacion v1.4.xlsx"
OUT_REPO = r"C:\Users\Usuario\OneDrive\Escritorio\TesisDG-ML\PROYECTO-TESIS-DG\docs\excel\P20261012_Historias de Usuario y Criterios de Validacion v1.4.xlsx"

wb = openpyxl.load_workbook(SRC)
ws = wb["HU"]


def set_escenario(row, criterio=None, contexto=None, evento=None, resultado=None):
    if criterio is not None:
        ws.cell(row=row, column=7).value = criterio
    if contexto is not None:
        ws.cell(row=row, column=8).value = contexto
    if evento is not None:
        ws.cell(row=row, column=9).value = evento
    if resultado is not None:
        ws.cell(row=row, column=10).value = resultado


# HU012 esc1 (CP028) -- ya no hay EM2022; el desglose de notas por área es
# la funcionalidad real (panel "Factores de riesgo" agregado en esta ronda).
set_escenario(30,
    resultado="Sistema muestra las 3 áreas académicas con el promedio más bajo (Bimestre 1-3) como los factores que más elevan su riesgo")

# HU022 esc1 (CP048) -- sin EM2022, el colegio propio solo segmenta por nivel.
set_escenario(50,
    resultado="Sistema agrupa a los estudiantes por nivel de riesgo (ALTO/MEDIO/BAJO); no segmenta todavía por un tipo más específico dentro de cada nivel")

# HU022 esc2 (CP049) -- rediseñado: ya no tiene sentido repetir "sin tipo
# específico" (esc1 ya lo dice); el escenario negativo real y honesto es
# el colegio sin modelo propio entrenado todavía.
set_escenario(51,
    criterio="Colegio sin modelo propio",
    contexto="Colegio sin modelo propio entrenado todavía",
    evento="Intenta ver la segmentación de estudiantes",
    resultado="Sistema muestra el estado 'Sin modelo propio configurado' (con acceso directo a Datos para subir el Excel) en vez de cualquier segmentación")

# HU028 esc1 (CP060) -- sin EM2022, la integración real es Notas + Conducta
# del propio colegio combinadas en un solo modelo (dos archivos, un resultado).
set_escenario(62,
    contexto="Archivos Excel de notas y de conducta del colegio disponibles",
    evento="Sube ambos archivos para entrenar el modelo",
    resultado="Sistema combina notas y conducta en un mismo modelo, aunque el colegio las suba como archivos separados")

# HU034 esc1 (CP072) -- ya NO es una limitación: el modelo propio ahora sí
# expone importancia global de variables (Random Forest, agregado en esta
# ronda) -- este escenario pasa de "no disponible" a realmente disponible.
set_escenario(74,
    resultado="Sistema muestra un ranking global de variables (Random Forest) con las que más pesan en las predicciones de ese colegio")

wb.save(OUT_DOWNLOADS)
wb.save(OUT_REPO)
print("guardado:", OUT_DOWNLOADS)
print("guardado:", OUT_REPO)
