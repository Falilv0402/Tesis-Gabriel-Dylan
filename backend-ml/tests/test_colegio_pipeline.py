"""
Tests del pipeline de entrenamiento del modelo por colegio (modelo/colegio/):
validador de nombre, versión más reciente por (salón, bimestre), cambio de
lista entre bimestres y modos de pronóstico según los bimestres cargados.
Usan datos sintéticos (nada de archivos ni alumnos reales).

Ejecutar con:  pytest backend-ml/tests/test_colegio_pipeline.py -v
"""
import sys
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "modelo" / "colegio"))

import parse_excels as pe  # noqa: E402
import train_colegio_model as tm  # noqa: E402


# ─── Nombre del colegio ───────────────────────────────────────────────────────

@pytest.mark.parametrize("nombre,esperado", [
    ("colegio_831305_upload", False), ("uploads_831305", False), ("Colegio 831305", False),
    ("Colegio", False), ("IE 831305", False), ("", False), ("nan", False), (None, False),
    ("REPORTE CONSOLIDADO DE NOTAS", False),
    ("Cole Andres Avelino Trapiche", True), ("Joseph And Mery", True), ("Colegio San Juan", True),
])
def test_nombre_colegio_valido(nombre, esperado):
    assert pe.nombre_colegio_valido(nombre) is esperado


def test_resolver_nombre_hereda_el_del_modelo_anterior(tmp_path):
    out = tmp_path / "colegio_1.pkl"
    joblib.dump({"nombre_colegio": "Colegio Bueno"}, out)
    assert tm._resolver_nombre_colegio("colegio_1_upload", out, "1") == "Colegio Bueno"
    assert tm._resolver_nombre_colegio("Nombre Nuevo", out, "1") == "Nombre Nuevo"
    assert tm._resolver_nombre_colegio("", tmp_path / "otro.pkl", "7") == "Colegio 7"


# ─── Versión más reciente por (salón, bimestre) ──────────────────────────────

def _filas(archivo, mtime, bim, nota):
    return pd.DataFrame({
        "salon": "2B", "n_alumno": [1, 2], "nombre": ["Ana", "Luis"], "_bim": float(bim),
        "_archivo": archivo, "_mtime": mtime, f"b{bim}_matematica": nota,
    })


def test_gana_el_archivo_mas_reciente_del_mismo_salon_y_bimestre():
    df = pd.concat([_filas("viejo.xlsx", 100.0, 2, 10.0), _filas("nuevo.xlsx", 200.0, 2, 15.0),
                    _filas("b1.xlsx", 50.0, 1, 12.0)], ignore_index=True)
    avisos: list[str] = []
    out = pe._quedarse_con_version_mas_reciente(df, avisos)
    assert set(out["_archivo"]) == {"nuevo.xlsx", "b1.xlsx"}
    assert (out.loc[out["_archivo"] == "nuevo.xlsx", "b2_matematica"] == 15.0).all()
    assert len(avisos) == 1 and "2B B2" in avisos[0]


def test_sin_repetidos_no_avisa():
    df = pd.concat([_filas("a.xlsx", 1.0, 1, 12.0), _filas("b.xlsx", 2.0, 2, 13.0)], ignore_index=True)
    avisos: list[str] = []
    assert len(pe._quedarse_con_version_mas_reciente(df, avisos)) == len(df)
    assert avisos == []


# ─── Cambio de lista entre bimestres ─────────────────────────────────────────

def test_alumno_retirado_se_cruza_por_nombre_y_conserva_el_numero_del_ultimo_excel():
    b1 = pd.DataFrame({"salon": "2B", "n_alumno": [1, 2, 3], "nombre": ["Ana Ruiz", "Beto Paz", "Carla Soto"],
                       "_bim": 1.0, "_archivo": "b1.xlsx", "_mtime": 1.0, "b1_matematica": [10.0, 11.0, 12.0]})
    # En B2 se retira Beto: Carla pasa a ser el N.° 2
    b2 = pd.DataFrame({"salon": "2B", "n_alumno": [1, 2], "nombre": ["Ana Ruiz", "Carla Soto"],
                       "_bim": 2.0, "_archivo": "b2.xlsx", "_mtime": 2.0, "b2_matematica": [13.0, 14.0]})
    avisos: list[str] = []
    out = pe._unificar_numeracion_por_nombre(pd.concat([b1, b2], ignore_index=True), avisos)
    assert sorted(out["nombre"].unique()) == ["Ana Ruiz", "Carla Soto"]          # Beto descartado
    carla = out[out["nombre"] == "Carla Soto"]
    assert (carla["n_alumno"] == 2).all()                                        # N.° del último Excel
    assert carla["b1_matematica"].dropna().iloc[0] == 12.0                       # su nota de B1, no la de Beto
    assert avisos and "descartaron" in avisos[0]


def test_lista_estable_no_toca_nada():
    df = pd.DataFrame({"salon": "2B", "n_alumno": [1, 2, 1, 2], "nombre": ["Ana", "Luis", "Ana", "Luis"],
                       "_bim": [1.0, 1.0, 2.0, 2.0], "_archivo": ["a", "a", "b", "b"], "_mtime": [1.0, 1.0, 2.0, 2.0]})
    avisos: list[str] = []
    out = pe._unificar_numeracion_por_nombre(df, avisos)
    assert len(out) == 4 and avisos == []


# ─── Bimestres cargados y riesgo por bimestre ────────────────────────────────

def _df_bims(bims, n=8, nota=14.0):
    d = {}
    for b in bims:
        d[f"b{b}_matematica"] = [nota] * n
        d[f"b{b}_comunicacion"] = [nota] * n
    return pd.DataFrame(d)


def test_bimestres_cargados_cuenta_consecutivos_desde_b1():
    assert tm._bimestres_cargados(_df_bims([1, 2, 3])) == 3
    assert tm._bimestres_cargados(_df_bims([1, 2, 3, 4])) == 4
    assert tm._bimestres_cargados(_df_bims([1, 3])) == 1      # un hueco corta la cuenta
    assert tm._bimestres_cargados(pd.DataFrame({"pp_matematica": [14.0] * 8})) == 0


def test_riesgo_bimestre_umbral_13_en_matematica_o_comunicacion():
    df = pd.DataFrame({"b1_matematica": [13.0, 14.0, 14.0, np.nan], "b1_comunicacion": [15.0, 12.0, 15.0, np.nan]})
    r = tm._riesgo_bimestre(df, 1).tolist()
    assert r[:3] == [1.0, 1.0, 0.0] and np.isnan(r[3])


# ─── Entrenamiento de punta a punta con Excel sintéticos ─────────────────────

CURSOS = ["ÁLGEBRA", "COMUNICACIÓN", "CIENCIA Y TECNOLOGÍA", "CIENCIAS SOCIALES", "INGLÉS", "DANZA",
          "EDUCACIÓN FÍSICA", "COMPORTAMIENTO"]
ROMANO = {1: "I", 2: "II", 3: "III", 4: "IV"}


def _escribir_excel(carpeta: Path, salon: str, bim: int, notas: np.ndarray, nombres: list[str]):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CONSOLIDADO DE NOTAS"
    ws["A1"] = f"REPORTE CONSOLIDADO DE NOTAS - {ROMANO[bim]} BIMESTRE"
    ws["A2"] = salon
    ws.append([])
    ws.append(["N°", "APELLIDOS Y NOMBRES", *CURSOS])
    for i, nombre in enumerate(nombres, start=1):
        ws.append([i, nombre, *[int(x) for x in notas[i - 1]]])
    wb.save(carpeta / f"Reporte consolidado de notas ({salon[-1]}{bim}).xlsx")


def _colegio_sintetico(carpeta: Path, bimestres, n_por_salon=40, nota_base=13.5, seed=0):
    rng = np.random.default_rng(seed)
    for letra in ("A", "B"):
        nombres = [f"ALUMNO {letra}{i:02d}, PRUEBA" for i in range(n_por_salon)]
        capacidad = rng.normal(nota_base, 2.2, n_por_salon)
        for b in bimestres:
            notas = capacidad[:, None] + rng.normal(0, 0.8, (n_por_salon, len(CURSOS)))
            _escribir_excel(carpeta, f"SECUNDARIA - 2° - {letra}", b, np.clip(np.round(notas), 5, 20), nombres)


@pytest.fixture
def entrenar(tmp_path, monkeypatch):
    # train() guarda en <parents[1]>/model del archivo del módulo: se redirige a tmp_path
    (tmp_path / "modelo" / "colegio").mkdir(parents=True)
    monkeypatch.setattr(tm, "__file__", str(tmp_path / "modelo" / "colegio" / "train_colegio_model.py"))

    def _run(bimestres, **kw):
        carpeta = tmp_path / "excels"
        carpeta.mkdir(exist_ok=True)
        for f in carpeta.glob("*.xlsx"):
            f.unlink()
        _colegio_sintetico(carpeta, bimestres, **kw)
        tm.train(str(carpeta), "999001")
        return joblib.load(tmp_path / "modelo" / "model" / "colegio_999001.pkl")
    return _run


def test_tres_bimestres_pronostican_el_cuarto(entrenar):
    art = entrenar([1, 2, 3])
    m = art["metricas"]
    assert m["tipo_resultado"] == "pronostico"
    assert m["bimestres_cargados"] == [1, 2, 3] and m["bimestre_pronosticado"] == 4
    assert m["modo_prediccion"].startswith("predictivo")
    assert m["auc_cv"] is not None and m["auc_cv"] > 0.75
    preds = art["predicciones"]
    assert len(preds) == 80
    assert all(0.0 <= p["prob_riesgo"] <= 1.0 for p in preds)
    assert all(0 <= p["bimestres_en_riesgo"] <= 3 and p["bimestres_evaluados"] == 3 for p in preds)
    assert {p["nivel_riesgo"] for p in preds} <= {"ALTO", "MEDIO", "BAJO"}


def test_dos_bimestres_pronostican_el_tercero(entrenar):
    m = entrenar([1, 2])["metricas"]
    assert m["tipo_resultado"] == "pronostico" and m["bimestre_pronosticado"] == 3


def test_cuatro_bimestres_son_cierre_de_anio_sin_bimestre_que_pronosticar(entrenar):
    art = entrenar([1, 2, 3, 4])
    m = art["metricas"]
    assert m["tipo_resultado"] == "cierre_anual" and m["bimestre_pronosticado"] is None
    assert m["modo_prediccion"].startswith("cierre de año")
    assert all(p["bimestres_evaluados"] == 4 for p in art["predicciones"])


def test_un_solo_bimestre_es_descriptivo_y_lo_dice(entrenar):
    m = entrenar([1])["metricas"]
    assert m["tipo_resultado"] == "descriptivo" and m["bimestre_pronosticado"] is None
    assert m["modo_prediccion"].startswith("descriptivo")


def test_muestra_sin_alumnos_en_riesgo_no_revienta_y_queda_descriptiva(entrenar):
    # Todos con notas altas: nadie en riesgo → no hay con qué entrenar; antes caía en silencio al índice
    art = entrenar([1, 2, 3], nota_base=18.5)
    m = art["metricas"]
    assert m["tipo_resultado"] == "descriptivo"
    assert m["modo_prediccion"].startswith("descriptivo")
    assert m["auc_cv"] is None
