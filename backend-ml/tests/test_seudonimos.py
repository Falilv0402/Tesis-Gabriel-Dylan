"""
Tests de la seudonimización (opción D): ningún nombre real queda en el sistema.
Usan datos sintéticos (nada de archivos ni alumnos reales).

Ejecutar con:  pytest backend-ml/tests/test_seudonimos.py -v
"""
import json
import os
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "modelo" / "colegio"))

import anonimizar as an  # noqa: E402
import parse_excels as pe  # noqa: E402
import seudonimos as sd  # noqa: E402


# ─── Seudonimizador ───────────────────────────────────────────────────────────

def test_etiqueta_y_deteccion_de_seudonimos():
    assert sd.etiqueta(12) == "Estudiante 000012"
    assert sd.codigo_de_etiqueta("Estudiante 000012") == 12
    assert sd.es_seudonimo("Estudiante 000012") and sd.es_seudonimo("  Estudiante 123456 ")
    assert not sd.es_seudonimo("Ana Pérez") and not sd.es_seudonimo(None) and not sd.es_seudonimo("Alumno 3")


def test_el_mismo_alumno_recibe_el_mismo_codigo_aunque_cambie_como_se_escribe():
    s = sd.Seudonimizador(None)
    a = s.codigo("BERNAL MONTERO, FÁTIMA", "2B", 1)
    assert s.codigo("Bernal Montero Fatima", "2B", 17) == a          # otra grafía y otro N.° de lista: misma persona
    assert s.codigo("  bernal   montero,  fatima ", "2B", 1) == a


def test_distinto_alumno_o_distinto_salon_codigos_distintos():
    s = sd.Seudonimizador(None)
    codigos = {s.codigo("PÉREZ RUIZ, ANA", "2B", 1), s.codigo("PÉREZ RUIZ, ANA", "2A", 1), s.codigo("RUIZ PÉREZ, LUIS", "2B", 2)}
    assert len(codigos) == 3
    assert sorted(codigos) == [1, 2, 3]            # secuenciales, sin saltos


def test_alumno_sin_nombre_se_identifica_por_salon_y_numero_de_lista():
    s = sd.Seudonimizador(None)
    a = s.codigo(None, "2B", 3)
    assert s.codigo("", "2B", 3) == a and s.codigo(float("nan"), "2B", 3) == a and s.codigo("Alumno 3", "2B", 3) == a
    assert s.codigo(None, "2B", 4) != a and s.codigo(None, "2A", 3) != a


def test_un_seudonimo_que_ya_viene_en_el_archivo_se_respeta_y_no_se_reutiliza(tmp_path):
    s = sd.Seudonimizador(tmp_path / "m.json")
    with s.sesion():
        assert s.codigo("Estudiante 000050", "2B", 1) == 50
        assert s.codigo("Fulano de Tal", "2B", 2) == 51            # el siguiente código libre queda después del 50


def test_persistencia_estable_y_el_mapa_no_guarda_nombres(tmp_path):
    ruta = tmp_path / "colegio_1_seudonimos.json"
    nombres = [("GARCÍA LÓPEZ, MARÍA", "1A", 1), ("TORRES VEGA, JOSÉ", "1A", 2), ("QUISPE HUAMÁN, ROSA", "1B", 1)]
    s1 = sd.Seudonimizador(ruta)
    with s1.sesion():
        primeros = [s1.codigo(*x) for x in nombres]
    s2 = sd.Seudonimizador(ruta)                    # otro proceso / otra carga
    with s2.sesion():
        assert [s2.codigo(*x) for x in nombres] == primeros
        assert s2.codigo("NUEVO ALUMNO", "1B", 2) == 4
    contenido = ruta.read_text(encoding="utf-8").upper()
    for nombre, _, _ in nombres + [("NUEVO ALUMNO", "", 0)]:
        for trozo in nombre.replace(",", "").split():
            assert trozo not in contenido                           # ni el nombre ni sus apellidos
    assert sorted(json.loads(contenido)["MAPA"].values()) == [1, 2, 3, 4]


def test_la_clave_se_crea_junto_al_mapa_y_la_variable_de_entorno_manda(tmp_path, monkeypatch):
    monkeypatch.delenv("SATRA_PSEUDONYM_KEY", raising=False)
    sd.Seudonimizador(tmp_path / "m.json")
    assert (tmp_path / ".seudonimos.key").exists()
    # con otra clave, el mismo nombre produce otra huella (el mapa de un servidor no sirve en otro)
    a, b = sd.Seudonimizador(None, clave=b"clave-a"), sd.Seudonimizador(None, clave=b"clave-b")
    assert a._hash("2B", "ANA") != b._hash("2B", "ANA")
    monkeypatch.setenv("SATRA_PSEUDONYM_KEY", "desde-el-entorno")
    assert sd.Seudonimizador(tmp_path / "otro" / "m.json")._clave == b"desde-el-entorno"
    assert not (tmp_path / "otro" / ".seudonimos.key").exists()


def test_aplicar_trabaja_por_columnas():
    s = sd.Seudonimizador(None)
    out = s.aplicar(["2B", "2B", "2A"], [1, 2, 1], ["ANA PÉREZ", "LUIS RUIZ", "ANA PÉREZ"])
    assert out == ["Estudiante 000001", "Estudiante 000002", "Estudiante 000003"]


# ─── Anonimizar Excel ─────────────────────────────────────────────────────────

NOMBRES = ["BERNAL MONTERO, FATIMA", "BLAS FABIAN, ABDIEL", "BRETONECHE PURIZACA, SOFIA", "BUSTAMANTE GALVEZ, DAPHNE",
           "CAPIRE CARBAJAL, SURREY", "CARBAJAL RUIZ, MATEO"]
ESPECIALES = ["BERNAL", "FABIAN", "BRETONECHE", "BUSTAMANTE", "CAPIRE", "CARBAJAL"]   # apellidos sintéticos a buscar


def _reporte_consolidado(ruta: Path, salon="SECUNDARIA - 2° - B", bim="II", nombres=NOMBRES, nota=14):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CONSOLIDADO DE NOTAS"
    ws["A1"] = f"REPORTE CONSOLIDADO DE NOTAS - {bim} BIMESTRE"
    ws["A2"] = salon
    ws.append([])
    ws.append(["N°", "APELLIDOS Y NOMBRES", "ÁLGEBRA", "COMUNICACIÓN", "COMPORTAMIENTO"])
    for i, nombre in enumerate(nombres, start=1):
        ws.append([i, nombre, nota + (i % 3), nota - (i % 4), 15])
    wb.properties.creator = "Secretaría del colegio"
    wb.save(ruta)


def _texto_de_todas_las_celdas(ruta: Path) -> str:
    partes = []
    for df in pd.read_excel(ruta, sheet_name=None, header=None, dtype=str).values():
        partes.append(" ".join(df.fillna("").astype(str).to_numpy().ravel()))
    return " ".join(partes).upper()


def test_anonimizar_reporte_consolidado_quita_los_nombres_y_devuelve_el_libro(tmp_path):
    origen, destino = tmp_path / "Reporte consolidado de notas (1).xlsx", tmp_path / "anon.xlsx"
    _reporte_consolidado(origen)
    s = sd.Seudonimizador(tmp_path / "m.json")
    with s.sesion():
        libro = an.anonimizar_libro(origen, destino, s, "notas")

    texto = _texto_de_todas_las_celdas(destino)
    assert not any(apellido in texto for apellido in ESPECIALES)
    assert "ESTUDIANTE 000001" in texto and "ESTUDIANTE 000006" in texto
    assert len(libro) == len(NOMBRES)
    assert {r["nombre"] for r in libro} == set(NOMBRES)                       # los nombres salen SOLO en el libro
    assert all(r["salon"] == "2B" and r["archivo"] == origen.name for r in libro)
    assert [r["n_alumno"] for r in libro] == list(range(1, 7))
    assert pd.read_excel(destino, header=None, dtype=str).iloc[3, 1] == "APELLIDOS Y NOMBRES"   # encabezados intactos
    # el original no se tocó y los metadatos del libro (autor) no viajan a la copia
    assert "BERNAL" in _texto_de_todas_las_celdas(origen)
    import openpyxl
    assert openpyxl.load_workbook(destino).properties.creator != "Secretaría del colegio"


def test_el_parser_da_lo_mismo_antes_y_despues_de_anonimizar(tmp_path):
    crudo, anon = tmp_path / "crudo", tmp_path / "anon"
    crudo.mkdir(), anon.mkdir()
    for bim, nota in (("I", 12), ("II", 14), ("III", 13)):
        _reporte_consolidado(crudo / f"Reporte consolidado de notas ({bim}).xlsx", bim=bim, nota=nota)
    s = sd.Seudonimizador(tmp_path / "m.json")
    with s.sesion():
        for f in sorted(crudo.glob("*.xlsx")):
            an.anonimizar_libro(f, anon / f.name, s, "notas")
        df_crudo = pe.procesar_colegio(crudo, "1")
        df_anon = pe.procesar_colegio(anon, "1", s)

    a = df_crudo.sort_values(["salon", "n_alumno"]).reset_index(drop=True)
    b = df_anon.sort_values(["salon", "n_alumno"]).reset_index(drop=True)
    assert a.shape == b.shape
    for col in a.columns:
        if col in ("nombre", "nombre_colegio"):
            continue
        if pd.api.types.is_numeric_dtype(a[col]):
            assert np.allclose(a[col].astype(float), b[col].astype(float), equal_nan=True), col
        else:
            assert a[col].fillna("").astype(str).equals(b[col].fillna("").astype(str)), col
    assert b["nombre"].map(sd.es_seudonimo).all() and b["nombre"].nunique() == len(NOMBRES)


def test_red_de_seguridad_del_pipeline_seudonimiza_aunque_el_excel_traiga_nombres(tmp_path):
    carpeta = tmp_path / "colegio_1_upload"
    carpeta.mkdir()
    _reporte_consolidado(carpeta / "Reporte consolidado de notas (1).xlsx")
    s = sd.Seudonimizador(tmp_path / "m.json")
    with s.sesion():
        df = pe.procesar_colegio(carpeta, "1", s)
    assert df["nombre"].map(sd.es_seudonimo).all()
    assert not any(apellido in " ".join(df["nombre"]).upper() for apellido in ESPECIALES)


def test_la_celda_contigua_del_nombre_cubicol_se_anonimiza_pero_una_nota_literal_no(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S4A"
    for fila in (["COLEGIO X"], ["CUBICOL"], ["CONSOLIDADO DE NOTAS A"], [], ["Salón:", "", "CUARTO AÑO A"],
                 ["ASIGNATURAS", "", "", "BIMESTRE", "MATEMÁTICA"], ["", "", "", "", "Resuelve problemas"],
                 ["N°", "APELLIDOS Y NOMBRES", "", "", "C1"],                       # CUBICOL: encabezado contiguo vacío
                 [1, None, "QUISPE HUAMÁN, ROSA", "1°", "A"],                       # el nombre cae en la columna 2
                 [None, None, None, "2°", "B"],
                 [2, "TORRES VEGA, JOSÉ", None, "1°", "AD"]):                       # el nombre cae en la columna 1
        ws.append(fila)
    # hoja de otro formato: la columna contigua tiene encabezado (RELIGIÓN) → "AD" es una nota, no un nombre
    ws2 = wb.create_sheet("CONSOLIDADO DE NOTAS")
    for fila in (["REPORTE CONSOLIDADO DE NOTAS - I BIMESTRE"], ["SECUNDARIA - 1° - A"], [], ["N°", "APELLIDOS Y NOMBRES", "RELIGIÓN"],
                 [1, None, "AD"]):
        ws2.append(fila)
    origen, destino = tmp_path / "Cuarto año sec A - Notas.xlsx", tmp_path / "out.xlsx"
    wb.save(origen)
    s = sd.Seudonimizador(None)
    libro = an.anonimizar_libro(origen, destino, s, "notas")
    texto = _texto_de_todas_las_celdas(destino)
    assert "QUISPE" not in texto and "TORRES" not in texto
    assert {r["nombre"] for r in libro} == {"QUISPE HUAMÁN, ROSA", "TORRES VEGA, JOSÉ"}      # 'AD' no se tomó por un nombre
    assert "AD" in pd.read_excel(destino, sheet_name="CONSOLIDADO DE NOTAS", header=None, dtype=str).iloc[4, 2]


def test_conserva_la_fecha_del_archivo_original(tmp_path):
    origen, destino = tmp_path / "Notas.xlsx", tmp_path / "Notas_anon.xlsx"
    _reporte_consolidado(origen)
    antes = time.time() - 86_400 * 30
    os.utime(origen, (antes, antes))
    s = sd.Seudonimizador(None)
    an.anonimizar_libro(origen, destino, s, "notas")
    assert abs(destino.stat().st_mtime - antes) < 2          # el parser decide la "versión más reciente" por esta fecha


def test_un_archivo_danado_no_deja_nada_en_la_carpeta_del_colegio(tmp_path):
    carpeta = tmp_path / "colegio"
    carpeta.mkdir()
    roto = tmp_path / "Notas.xlsx"
    roto.write_bytes(b"esto no es un excel")
    s = sd.Seudonimizador(None)
    with pytest.raises(Exception):
        an.anonimizar_subida(roto, carpeta, "Notas.xlsx", s)
    assert list(carpeta.iterdir()) == []


def test_anonimizar_subida_usa_el_nombre_original_con_extension_xlsx(tmp_path):
    carpeta = tmp_path / "colegio"
    carpeta.mkdir()
    temporal = tmp_path / "subido.tmp"
    _reporte_consolidado(temporal)
    s = sd.Seudonimizador(tmp_path / "m.json")
    with s.sesion():
        destino, libro = an.anonimizar_subida(temporal, carpeta, "Reporte consolidado de notas.xls", s)
    assert destino == carpeta / "Reporte consolidado de notas.xlsx" and destino.exists()
    assert len(libro) == len(NOMBRES)
    assert [p.name for p in carpeta.iterdir()] == ["Reporte consolidado de notas.xlsx"]    # ni temporales ni el original


def test_volver_a_subir_el_mismo_excel_da_los_mismos_codigos(tmp_path):
    """Si el colegio pierde el libro de claves, lo regenera subiendo otra vez el mismo archivo."""
    origen = tmp_path / "Notas.xlsx"
    _reporte_consolidado(origen)
    s = sd.Seudonimizador(tmp_path / "m.json")
    with s.sesion():
        libro1 = an.anonimizar_libro(origen, tmp_path / "a.xlsx", s, "notas")
    s2 = sd.Seudonimizador(tmp_path / "m.json")
    with s2.sesion():
        libro2 = an.anonimizar_libro(origen, tmp_path / "b.xlsx", s2, "notas")
    assert [(r["codigo"], r["nombre"]) for r in libro1] == [(r["codigo"], r["nombre"]) for r in libro2]


# ─── Endpoint /procesar de punta a punta ──────────────────────────────────────

def _colegio_sintetico(carpeta: Path, n_por_salon=40, seed=3):
    """3 bimestres × 2 salones de 40 alumnos, con nombres sintéticos reconocibles ('ALUMNO A00, PRUEBA')."""
    import openpyxl
    rng = np.random.default_rng(seed)
    cursos = ["ÁLGEBRA", "COMUNICACIÓN", "CIENCIA Y TECNOLOGÍA", "CIENCIAS SOCIALES", "INGLÉS", "DANZA", "EDUCACIÓN FÍSICA", "COMPORTAMIENTO"]
    for letra in ("A", "B"):
        nombres = [f"ALUMNO {letra}{i:02d}, PRUEBA" for i in range(n_por_salon)]
        capacidad = rng.normal(13.5, 2.2, n_por_salon)
        for bim, romano in ((1, "I"), (2, "II"), (3, "III")):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "CONSOLIDADO DE NOTAS"
            ws["A1"] = f"REPORTE CONSOLIDADO DE NOTAS - {romano} BIMESTRE"
            ws["A2"] = f"SECUNDARIA - 2° - {letra}"
            ws.append([])
            ws.append(["N°", "APELLIDOS Y NOMBRES", *cursos])
            for i, nombre in enumerate(nombres, start=1):
                notas = np.clip(np.round(capacidad[i - 1] + rng.normal(0, 0.8, len(cursos))), 5, 20)
                ws.append([i, nombre, *[int(x) for x in notas]])
            wb.save(carpeta / f"Reporte consolidado de notas ({letra}{bim}).xlsx")


def test_endpoint_procesar_no_deja_nombres_en_ningun_lado_y_entrega_el_libro(tmp_path, monkeypatch):
    import shutil
    import tempfile

    from starlette.testclient import TestClient

    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from app.main import app
    from app.api import colegio_propio as cp

    # Todo el flujo (carpeta del colegio, modelo, mapa) en una copia temporal: nada toca el repo
    raiz = tmp_path / "modelo"
    shutil.copytree(Path(cp.COLEGIO_SCRIPTS), raiz / "colegio", ignore=shutil.ignore_patterns("__pycache__"))
    (raiz / "model").mkdir()
    (raiz / "data").mkdir()
    monkeypatch.setattr(cp, "COLEGIO_SCRIPTS", raiz / "colegio")
    monkeypatch.setattr(cp, "MODEL_DIR", raiz / "model")
    monkeypatch.setattr(cp, "DATA_DIR", raiz / "data")

    async def _sin_supabase(*a, **k):
        return 0
    monkeypatch.setattr(cp, "_alertar_nuevos_alto", _sin_supabase)
    monkeypatch.setattr(cp, "_registrar_version_modelo", _sin_supabase)
    app.dependency_overrides[cp.require_admin_de_colegio] = lambda: {"codigo_ie": "999780", "token": "t"}

    origen = tmp_path / "originales"
    origen.mkdir()
    _colegio_sintetico(origen)
    archivos = sorted(origen.glob("*.xlsx"))
    temporales_antes = set(Path(tempfile.gettempdir()).glob("satra_subida_*"))

    def subir():
        files = [("notas_files", (f.name, f.read_bytes(), "application/octet-stream")) for f in archivos]
        files.append(("notas_files", ("Notas viejas.pdf", b"%PDF", "application/pdf")))      # uno inválido: solo ese se rechaza
        return TestClient(app).post("/v1/colegio/999780/procesar", files=files, data={"nombre_colegio": "Colegio de Prueba"})

    try:
        r1 = subir()
        assert r1.status_code == 200, r1.text
        j1 = r1.json()
        assert r1.headers["cache-control"] == "no-store"
        assert j1["archivos_procesados"] == 6 and j1["archivos_rechazados"] == ["Notas viejas.pdf: Formato no soportado: solo .xlsx o .xls"]

        # libro de claves: 80 alumnos, con sus nombres reales (es lo único que los lleva)
        libro = j1["libro_claves"]
        assert len(libro) == 80 and len({r["codigo"] for r in libro}) == 80
        assert all(r["nombre"].endswith("PRUEBA") and r["codigo"].startswith("Estudiante ") for r in libro)
        assert j1["metricas"]["nombre_colegio"] == "Colegio de Prueba"

        # en disco no queda ningún nombre: ni en la carpeta del colegio, ni en el modelo, ni en el mapa
        carpeta = raiz / "data" / "colegio_999780_upload"
        assert sorted(p.suffix for p in carpeta.glob("*.xls*")) == [".xlsx"] * 6
        for f in carpeta.iterdir():
            contenido = _texto_de_todas_las_celdas(f) if f.suffix == ".xlsx" else f.read_text(encoding="utf-8", errors="ignore").upper()
            assert "ALUMNO A0" not in contenido and "PRUEBA" not in contenido, f.name
        # bytes del ARCHIVO (no una re-serialización: en este proceso hay nombres sintéticos en la memoria)
        assert b"ALUMNO A0" not in (raiz / "model" / "colegio_999780.pkl").read_bytes()
        assert "PRUEBA" not in (raiz / "model" / "colegio_999780_seudonimos.json").read_text(encoding="utf-8").upper()
        assert set(Path(tempfile.gettempdir()).glob("satra_subida_*")) == temporales_antes      # el temporal con los originales se borró

        # subir otra vez los mismos Excel → los mismos códigos (así el colegio puede regenerar su libro)
        j2 = subir().json()
        assert {(r["nombre"], r["codigo"]) for r in j2["libro_claves"]} == {(r["nombre"], r["codigo"]) for r in libro}
    finally:
        app.dependency_overrides.clear()


# ─── Migración de carpetas que ya existen ─────────────────────────────────────

def test_migracion_en_el_lugar_conserva_fechas_borra_el_csv_y_verifica(tmp_path):
    import csv

    import migrar_a_seudonimos as mg

    carpeta = tmp_path / "colegio_1_upload"
    carpeta.mkdir()
    fechas = {}
    for k, (bim, nota) in enumerate((("I", 12), ("II", 14), ("III", 13)), start=1):
        f = carpeta / f"Reporte consolidado de notas ({bim}).xlsx"
        _reporte_consolidado(f, bim=bim, nota=nota)
        fechas[f.name] = time.time() - 86_400 * 10 * k
        os.utime(f, (fechas[f.name], fechas[f.name]))
    (carpeta / "colegio_1_procesado.csv").write_text("nombre,riesgo\n" + "\n".join(f"{n},1" for n in NOMBRES), encoding="utf-8")

    libro = tmp_path / "salida" / "libro.csv"
    resumen = mg.migrar(carpeta, "1", tmp_path / "m.json", libro)

    assert resumen["archivos"] == 3 and resumen["alumnos_en_el_libro"] == len(NOMBRES) and resumen["csv_borrados"] == ["colegio_1_procesado.csv"]
    assert sorted(p.name for p in carpeta.iterdir()) == sorted(fechas)                   # mismos archivos, sin el CSV
    for nombre, fecha in fechas.items():
        assert abs((carpeta / nombre).stat().st_mtime - fecha) < 2                       # la fecha original se conserva
    assert not any(a in _texto_de_todas_las_celdas(f) for f in carpeta.glob("*.xlsx") for a in ESPECIALES)
    assert mg.verificar(carpeta, libro) == []

    filas = list(csv.DictReader(open(libro, encoding="utf-8-sig", newline="")))
    assert [f["codigo"] for f in filas] == [sd.etiqueta(i) for i in range(1, 7)]
    assert {f["nombre"] for f in filas} == set(NOMBRES) and {f["salon"] for f in filas} == {"2B"}

    # si se cuela un archivo crudo, la verificación lo detecta
    _reporte_consolidado(carpeta / "Reporte consolidado de notas (IV).xlsx", bim="IV")
    assert mg.verificar(carpeta, libro)


def test_el_sql_de_limpieza_de_supabase_solo_lleva_seudonimos(tmp_path):
    import joblib

    import generar_sql_seudonimos as gs

    preds = [{"codigo_ie": "0831305", "salon": "2B", "n_alumno": 1, "nombre": "Estudiante 000001"},
             {"codigo_ie": "0831305", "salon": "2B", "n_alumno": 2, "nombre": "Estudiante 000002"}]
    joblib.dump({"predicciones": preds}, tmp_path / "colegio_831305.pkl")
    sql = gs.generar_sql(tmp_path, ["831305"])
    assert "('0831305-2B-1', 'Estudiante 000001')" in sql and "('0831305-2B-2', 'Estudiante 000002')" in sql
    assert "ltrim(x.estudiante_id, '0') = ltrim(m.id, '0')" in sql                  # tolera el cero inicial del código
    assert "UPDATE public.notificaciones" in sql and "UPDATE public.plan_estado" in sql

    # si el modelo todavía trae nombres reales, se niega a generar (hay que migrar y reentrenar antes)
    preds[1]["nombre"] = "BERNAL MONTERO, FATIMA"
    joblib.dump({"predicciones": preds}, tmp_path / "colegio_831305.pkl")
    with pytest.raises(SystemExit):
        gs.generar_sql(tmp_path, ["831305"])


# ─── Un archivo solo se guarda si se puede garantizar que no lleva nombres ────

def _libro_con_hojas(ruta: Path, hojas: dict[str, list[list]]):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        for fila in filas:
            ws.append(fila)
    wb.save(ruta)


def test_un_archivo_sin_la_columna_de_nombres_se_rechaza_y_no_se_guarda(tmp_path):
    carpeta = tmp_path / "colegio"
    carpeta.mkdir()
    origen = tmp_path / "Notas raras.xlsx"
    # otro encabezado ("ALUMNO"): el anonimizador no puede ubicar los nombres → no debe guardarlo tal cual
    _libro_con_hojas(origen, {"2B": [["N°", "ALUMNO", "MATEMÁTICA"], [1, "BERNAL MONTERO, FATIMA", 14], [2, "BLAS FABIAN, ABDIEL", 11]]})
    s = sd.Seudonimizador(None)
    with pytest.raises(ValueError, match="APELLIDOS Y NOMBRES"):
        an.anonimizar_subida(origen, carpeta, "Notas raras.xlsx", s)
    assert list(carpeta.iterdir()) == []


def test_una_hoja_worksheet_sin_columna_se_omite_pero_el_resto_se_guarda(tmp_path):
    origen, destino = tmp_path / "Notas.xlsx", tmp_path / "anon.xlsx"
    _libro_con_hojas(origen, {
        "CONSOLIDADO DE NOTAS": [["REPORTE CONSOLIDADO DE NOTAS - I BIMESTRE"], ["SECUNDARIA - 2° - B"], [],
                                 ["N°", "APELLIDOS Y NOMBRES", "ÁLGEBRA"], [1, "BERNAL MONTERO, FATIMA", 14]],
        "Worksheet 1": [["BLAS FABIAN, ABDIEL"], ["CAPIRE CARBAJAL, SURREY"]],     # sobrante de CUBICOL: el parser lo ignora
        "Hoja vacía": [],
    })
    s = sd.Seudonimizador(None)
    an.anonimizar_libro(origen, destino, s, "notas")
    hojas = pd.read_excel(destino, sheet_name=None, header=None, dtype=str)
    assert "Worksheet 1" not in hojas and "CONSOLIDADO DE NOTAS" in hojas
    texto = _texto_de_todas_las_celdas(destino)
    assert "BERNAL" not in texto and "BLAS" not in texto and "CAPIRE" not in texto
